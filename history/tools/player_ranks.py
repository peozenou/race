"""Year-by-year ATP ranking of every player ever rostered in THE RACE (2019-2026).

Usage: python3 -I history/tools/player_ranks.py SRC_DIR

SRC_DIR is the same private source folder as build_history.py (sackmann/ + season workbooks).
Reads history/data/rosters.json and writes history/data/ranks.json: one row per person and season,
whether or not a league team held the player that season, so departed players can be followed.

Sources by season:
  2019-2023  Jeff Sackmann tennis_atp, last ranking date of the year (December)
  2024       THE RACE 2024 workbook, tab 20241111w38 (ATP table, 11 Nov 2024, top 500)
  2025       THE RACE 2025 workbook, tab 20251117w35 (ATP table, 17 Nov 2025, top 500)
  2026       THE RACE 2026 workbook, latest dated tab (ATP table, season in progress)
"""
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from common import Identity, fold, num  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
SRC = Path(sys.argv[1])
IDS = Identity(SRC / "sackmann" / "atp_players.csv")
TE_TABS = {2024: ("THE_RACE_2024.xlsx", "20241111w38"), 2025: ("THE_RACE_2025.xlsx", "20251117w35")}


def sackmann_year_end(season):
    fn = "atp_rankings_2010s.csv" if season <= 2019 else "atp_rankings_20s.csv"
    best = {}
    with open(SRC / "sackmann" / fn, newline="") as f:
        for r in csv.DictReader(f):
            if r["ranking_date"][:4] != str(season):
                continue
            pid = int(r["player"])
            if pid not in best or r["ranking_date"] > best[pid][0]:
                best[pid] = (r["ranking_date"], int(r["rank"]), num(r["points"]))
    date = max(v[0] for v in best.values())
    return {pid: (rk, pts) for pid, (d, rk, pts) in best.items() if d == date}, f"{date[:4]}-{date[4:6]}-{date[6:]}"


def te_atp(fname, tab):
    ws = openpyxl.load_workbook(SRC / fname, data_only=True, read_only=True)[tab]
    by_id, by_name = {}, {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        if len(r) > 12 and r[10] and num(r[8]) is not None:
            rk, pts = num(r[8]), num(r[12])
            rec = IDS.resolve(r[10])
            if rec and rec["id"]:
                by_id.setdefault(rec["id"], (rk, pts))
            by_name.setdefault(frozenset(fold(r[10]).split()), (rk, pts))
    d = tab[:8]
    return by_id, by_name, f"{d[:4]}-{d[4:6]}-{d[6:]}"


def main():
    rosters = json.loads((ROOT / "history" / "data" / "rosters.json").read_text())
    people = {}
    for r in rosters:
        key = f"atp:{r['atp_id']}" if r["atp_id"] else f"name:{r['player']}"
        p = people.setdefault(key, {"person_key": key, "atp_id": r["atp_id"], "player": r["player"], "full": r["full"], "teams": {}})
        p["teams"].setdefault(r["season"], []).append(r["team"])
    tabs = dict(TE_TABS)
    wb26 = openpyxl.load_workbook(SRC / "THE_RACE_2026.xlsx", read_only=True)
    tabs[2026] = ("THE_RACE_2026.xlsx", max((s for s in wb26.sheetnames if re.match(r"^\d{8}w\d+$", s)), key=lambda s: s[:8]))
    out = []
    for season in range(2019, 2027):
        if season <= 2023:
            ids, date = sackmann_year_end(season)
            names, src = {}, "Jeff Sackmann tennis_atp, classement de fin d'année"
        else:
            ids, names, date = te_atp(*tabs[season])
            src = f"Tennis Explorer via THE RACE {season} (onglet {tabs[season][1]}, top 500)"
        for p in people.values():
            hit = ids.get(p["atp_id"]) if p["atp_id"] else None
            if hit is None and names:
                hit = names.get(frozenset(fold(p["full"]).split()))
            teams = p["teams"].get(season, [])
            out.append({"season": season, "person_key": p["person_key"], "atp_id": p["atp_id"], "player": p["player"],
                        "full": p["full"], "atp_rank": hit[0] if hit else None, "atp_points": hit[1] if hit else None,
                        "ranking_date": date, "in_progress": season == 2026,
                        "league_team": "/".join(dict.fromkeys(teams)) or None, "source": src})
    (ROOT / "history" / "data" / "ranks.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=0, separators=(",", ": ")) + "\n", encoding="utf-8")
    miss = sorted({(r["season"], r["player"]) for r in out if r["atp_rank"] is None and r["league_team"]})
    print(f"ranks.json: {len(out)} rows for {len(people)} players; rostered seasons without a rank: {miss}")


if __name__ == "__main__":
    main()
